"""
Persistance locale du plan (fichier Excel) et génération des exports.

Fichier de données : `data/plan_marathon.xlsx` (dossier modifiable via la variable
d'environnement PLAN_DATA_DIR). Il contient 3 onglets lisibles dans Excel :
    * parametres : clé / valeur (nom du plan, date de course, objectif par défaut)
    * semaines   : week_id, position, phase
    * seances    : une ligne par séance

Chaque sauvegarde est atomique (écriture dans un fichier temporaire puis
remplacement) et la version précédente est conservée en `.bak.xlsx`.
"""

from __future__ import annotations

import io
import os
import shutil
import tempfile
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import DISPLAY_NAMES, PACE_LABELS, SESSION_COLUMNS, WEEK_COLUMNS
from .paces import Pace, format_hms, format_minutes
from .plan import Plan
from .template import TEMPLATE_SETTINGS, TEMPLATE_WEEKS

APP_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = Path(os.environ.get("PLAN_DATA_DIR", APP_DIR / "data"))
DATA_FILE = DATA_DIR / "plan_marathon.xlsx"

SHEET_SETTINGS, SHEET_WEEKS, SHEET_SESSIONS = "parametres", "semaines", "seances"


# ----------------------------------------------------------------------------
# Plan par défaut
# ----------------------------------------------------------------------------
def default_plan() -> Plan:
    """Plan issu du fichier Excel d'origine (S0 -> S7)."""
    return Plan.from_week_list(TEMPLATE_WEEKS, TEMPLATE_SETTINGS)


# ----------------------------------------------------------------------------
# Sérialisation (format interne)
# ----------------------------------------------------------------------------
def plan_to_bytes(plan: Plan) -> bytes:
    """Plan -> classeur Excel (format interne de sauvegarde)."""
    plan = plan.normalized()
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(list(plan.settings.items()), columns=["cle", "valeur"]).to_excel(
            writer, sheet_name=SHEET_SETTINGS, index=False)
        plan.weeks[WEEK_COLUMNS].to_excel(writer, sheet_name=SHEET_WEEKS, index=False)
        plan.sessions[SESSION_COLUMNS].to_excel(writer, sheet_name=SHEET_SESSIONS, index=False)
        for ws in writer.book.worksheets:
            for i, width in enumerate([12, 10, 10, 10, 12, 34, 10, 12, 40], start=1):
                ws.column_dimensions[get_column_letter(i)].width = width
    return buffer.getvalue()


def plan_from_bytes(content: bytes) -> Plan:
    """Classeur Excel (format interne) -> Plan. Lève ValueError si le fichier est invalide."""
    try:
        sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, dtype=object, engine="openpyxl")
    except Exception as exc:  # fichier corrompu, mauvais format…
        raise ValueError(f"Fichier illisible : {exc}") from exc
    missing = {SHEET_SETTINGS, SHEET_WEEKS, SHEET_SESSIONS} - set(sheets)
    if missing:
        raise ValueError(f"Onglet(s) manquant(s) : {', '.join(sorted(missing))}. "
                         "Utilise un fichier exporté via « Sauvegarde des données ».")

    settings_df = sheets[SHEET_SETTINGS]
    settings = {str(k): ("" if pd.isna(v) else v) for k, v in zip(settings_df["cle"], settings_df["valeur"])}
    if isinstance(settings.get("race_date"), (datetime, date)):
        settings["race_date"] = settings["race_date"].strftime("%Y-%m-%d")
    settings = {k: str(v) for k, v in settings.items()}

    weeks = sheets[SHEET_WEEKS]
    sessions = sheets[SHEET_SESSIONS]
    for col in WEEK_COLUMNS:
        if col not in weeks.columns:
            raise ValueError(f"Colonne « {col} » absente de l'onglet {SHEET_WEEKS}.")
    for col in SESSION_COLUMNS:
        if col not in sessions.columns:
            sessions[col] = None
    weeks = weeks[WEEK_COLUMNS].copy()
    weeks["position"] = pd.to_numeric(weeks["position"], errors="coerce").fillna(0).astype(int)
    sessions = sessions[SESSION_COLUMNS].copy()
    sessions["ordre"] = pd.to_numeric(sessions["ordre"], errors="coerce").fillna(0).astype(int)
    return Plan(weeks, sessions, settings).normalized()


# ----------------------------------------------------------------------------
# Fichier local
# ----------------------------------------------------------------------------
def ensure_data_file(path: Path = DATA_FILE) -> Path:
    """Crée le fichier de données à partir du modèle Excel s'il n'existe pas encore."""
    if not path.exists():
        save_plan(default_plan(), path)
    return path


def load_plan(path: Path = DATA_FILE) -> Plan:
    """Charge le plan depuis le fichier local (créé à partir du modèle si absent).

    Si le fichier est corrompu, on tente la sauvegarde `.bak`, puis le modèle.
    """
    ensure_data_file(path)
    try:
        return plan_from_bytes(path.read_bytes())
    except ValueError:
        backup = _backup_path(path)
        if backup.exists():
            try:
                return plan_from_bytes(backup.read_bytes())
            except ValueError:
                pass
        return default_plan()


def save_plan(plan: Plan, path: Path = DATA_FILE) -> datetime:
    """Sauvegarde atomique du plan ; conserve la version précédente en .bak.xlsx."""
    path.parent.mkdir(parents=True, exist_ok=True)
    content = plan_to_bytes(plan)
    if path.exists():
        shutil.copy2(path, _backup_path(path))
    fd, tmp_name = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(content)
        os.chmod(tmp_name, 0o644)  # mkstemp crée le fichier en 0600
        os.replace(tmp_name, path)
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)
    return datetime.now()


def _backup_path(path: Path) -> Path:
    return path.with_name(path.stem + ".bak.xlsx")


# ----------------------------------------------------------------------------
# Exports pour l'utilisateur (plan généré avec ses allures)
# ----------------------------------------------------------------------------
EXPORT_COLUMNS = ["semaine", "debut", "fin", "phase", "seance", "km_env", "duree_env", "contenu",
                  "km_spe", "allure", "allure_txt", "allure_dec", "vitesse", "temps_spe_min", "remarque"]


def export_table(table: pd.DataFrame) -> pd.DataFrame:
    """Prépare le tableau du plan pour l'export (colonnes et libellés français)."""
    out = table[EXPORT_COLUMNS].copy()
    out["allure"] = [PACE_LABELS.get(c, "") if c else "" for c in out["allure"]]
    out["temps_spe_min"] = out["temps_spe_min"].round(0)
    out["debut"] = pd.to_datetime(out["debut"]).dt.date
    out["fin"] = pd.to_datetime(out["fin"]).dt.date
    return out.rename(columns=DISPLAY_NAMES)


def paces_table(paces: dict[str, Pace]) -> pd.DataFrame:
    """Tableau des allures de référence (équivalent du bloc K5:Q10 de l'Excel)."""
    rows = []
    for p in paces.values():
        rows.append({
            "Type d'allure": p.label,
            "Correspond à": p.description,
            "Formule (Excel)": p.excel_formula,
            "Vitesse (km/h)": round(p.kmh, 1),
            "Allure (min,cent./km)": round(p.min_per_km, 2),
            "Allure (min/km)": p.pace_str,
            "Temps de référence": (f"{p.ref_label} en {format_minutes(p.ref_time_minutes)}"
                                   if p.ref_time_minutes else ""),
        })
    return pd.DataFrame(rows)


def export_excel(table: pd.DataFrame, paces: dict[str, Pace], objective_seconds: int, plan: Plan) -> bytes:
    """Export Excel mis en forme : onglets Plan, Allures et Objectif."""
    plan_df = export_table(table)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        plan_df.to_excel(writer, sheet_name="Plan", index=False)
        paces_table(paces).to_excel(writer, sheet_name="Allures", index=False)
        pd.DataFrame([
            ("Plan", plan.settings.get("plan_name", "")),
            ("Objectif marathon", format_hms(objective_seconds)),
            ("Date de course", plan.race_date.strftime("%d/%m/%Y")),
            ("Nombre de semaines", plan.n_weeks),
            ("Généré le", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ], columns=["Paramètre", "Valeur"]).to_excel(writer, sheet_name="Objectif", index=False)

        header_fill = PatternFill("solid", fgColor="1F3A5F")
        week_fill = PatternFill("solid", fgColor="EEF2F7")
        for ws in writer.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for col_cells in ws.columns:
                length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(10, length + 2), 45)

        # Mise en forme de l'onglet Plan : dates, bandes alternées par semaine.
        ws = writer.book["Plan"]
        headers = [c.value for c in ws[1]]
        date_cols = [headers.index(DISPLAY_NAMES[k]) + 1 for k in ("debut", "fin")]
        week_col = headers.index(DISPLAY_NAMES["semaine"]) + 1
        shade, previous = False, None
        for row in ws.iter_rows(min_row=2):
            week = row[week_col - 1].value
            if week != previous:
                shade, previous = not shade, week
            for cell in row:
                if shade:
                    cell.fill = week_fill
                if cell.column in date_cols:
                    cell.number_format = "DD/MM/YYYY"
    return buffer.getvalue()


def export_csv(table: pd.DataFrame) -> bytes:
    """Export CSV compatible Excel français (séparateur « ; », virgule décimale, UTF-8 BOM)."""
    df = export_table(table)
    for col in (DISPLAY_NAMES["debut"], DISPLAY_NAMES["fin"]):
        df[col] = pd.to_datetime(df[col]).dt.strftime("%d/%m/%Y")
    return df.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")
